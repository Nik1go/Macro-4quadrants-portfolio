"""Backtest d'un strip ICE Brent fractionnaire contre les factures CNR.

Le classeur Bloomberg contient des rangs mobiles CO1..CO12. Chaque jour, on
retrouve le rang de l'échéance ICE réellement détenue avant de lire son prix.
Le modèle ne présente jamais cette position synthétique comme un contrat MBC.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO
from math import isfinite
from pathlib import Path
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl import load_workbook
from pandas.tseries.offsets import BDay

from .study import DATA_DIR


CONTRACT_BARRELS_ICE = 1_000
CONTRACT_BARRELS_MINI = 100
CNR_PATH = DATA_DIR / "cnr_gazole_professionnel_monthly.csv"
ECB_URL = "https://data-api.ecb.europa.eu/service/data/EXR/D.USD.EUR.SP00.A"


@dataclass(frozen=True)
class IceHistory:
    prices: pd.DataFrame  # index : jour de cotation ; colonnes : rangs 1..12
    volumes: pd.DataFrame
    open_interest: pd.DataFrame
    contracts: pd.DataFrame  # index : mois de livraison ; expiry, ticker


@dataclass(frozen=True)
class BacktestResult:
    monthly: pd.DataFrame
    daily: pd.DataFrame


def _number(value: object, label: str) -> float:
    if value is None or value == "":
        raise ValueError(f"Valeur manquante : {label}.")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Valeur non numérique : {label}.") from exc
    if not isfinite(result):
        raise ValueError(f"Valeur non finie : {label}.")
    return result


def _ice_expiry(theoretical: pd.Timestamp) -> pd.Timestamp:
    """ICE avance l'expiration si elle tombe la veille ouvrée de Noël ou du Nouvel An."""
    christmas_eve = pd.Timestamp(theoretical.year, 12, 25) - BDay(1)
    new_year_eve = pd.Timestamp(theoretical.year + 1, 1, 1) - BDay(1)
    if theoretical in (christmas_eve, new_year_eve):
        return (theoretical - BDay(1)).normalize()
    return theoretical


def load_ice_history(source: bytes | str | Path) -> IceHistory:
    """Lit l'export personnel Bloomberg sans le recopier dans le dépôt."""
    workbook_source = BytesIO(source) if isinstance(source, bytes) else source
    workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    try:
        generics = next(
            (sheet for sheet in workbook.worksheets if sheet.title.casefold().startswith("génériques")),
            None,
        )
        if generics is None or "Contrats" not in workbook.sheetnames:
            raise ValueError("Le classeur doit contenir les onglets Génériques et Contrats.")

        ticker_header = next(generics.iter_rows(min_row=5, max_row=5, max_col=37, values_only=True))
        field_header = next(generics.iter_rows(min_row=6, max_row=6, max_col=37, values_only=True))
        for rank in range(1, 13):
            column = 1 + 3 * (rank - 1)
            if str(ticker_header[column]).strip() != f"CO{rank} Comdty":
                raise ValueError(f"Ticker CO{rank} introuvable dans Génériques.")
            if tuple(field_header[column : column + 3]) != ("PX_LAST", "PX_VOLUME", "OPEN_INT"):
                raise ValueError(f"Champs de CO{rank} inattendus dans Génériques.")

        dates: list[pd.Timestamp] = []
        price_rows: list[list[float]] = []
        volume_rows: list[list[float | None]] = []
        oi_rows: list[list[float | None]] = []
        for row in generics.iter_rows(min_row=7, max_col=37, values_only=True):
            if not isinstance(row[0], (date, pd.Timestamp)):
                continue
            quote_date = pd.Timestamp(row[0]).normalize()
            prices = [_number(row[1 + 3 * i], f"CO{i + 1} au {quote_date.date()}") for i in range(12)]
            if any(price <= 0 for price in prices):
                raise ValueError(f"Prix Brent non positif au {quote_date.date()}.")
            dates.append(quote_date)
            price_rows.append(prices)
            volume_rows.append([row[2 + 3 * i] for i in range(12)])
            oi_rows.append([row[3 + 3 * i] for i in range(12)])
        if not dates or len(dates) != len(set(dates)):
            raise ValueError("Les dates Brent sont vides ou répétées.")
        columns = list(range(1, 13))
        prices = pd.DataFrame(price_rows, index=pd.DatetimeIndex(dates), columns=columns).sort_index()
        volumes = pd.DataFrame(volume_rows, index=pd.DatetimeIndex(dates), columns=columns).sort_index()
        open_interest = pd.DataFrame(oi_rows, index=pd.DatetimeIndex(dates), columns=columns).sort_index()

        contract_rows: list[dict[str, object]] = []
        for row in workbook["Contrats"].iter_rows(min_row=4, max_col=5, values_only=True):
            if not isinstance(row[0], (date, pd.Timestamp)):
                continue
            expiry = _ice_expiry(pd.Timestamp(row[2]).normalize())
            if prices.index.min() <= expiry <= prices.index.max() and expiry not in prices.index:
                quoted_days = prices.index[(prices.index <= expiry) & (prices.index.to_period("M") == expiry.to_period("M"))]
                if len(quoted_days):
                    expiry = quoted_days[-1]
            contract_rows.append(
                {
                    "delivery": pd.Timestamp(row[0]).to_period("M"),
                    "expiry": expiry,
                    "ticker": str(row[4]).strip(),
                }
            )
        contracts = pd.DataFrame(contract_rows)
        if contracts.empty or contracts["delivery"].duplicated().any():
            raise ValueError("Liste des contrats ICE vide ou avec des mois répétés.")
        contracts = contracts.sort_values("delivery").set_index("delivery")
        if contracts["ticker"].eq("").any() or contracts["expiry"].isna().any():
            raise ValueError("Ticker ou date d'expiration ICE manquant.")
        return IceHistory(prices, volumes, open_interest, contracts)
    finally:
        workbook.close()


def load_cnr_monthly(path: Path = CNR_PATH) -> pd.Series:
    data = pd.read_csv(path)
    if not {"date", "cnr"}.issubset(data.columns):
        raise ValueError("Le fichier CNR doit contenir date et cnr.")
    periods = pd.to_datetime(data["date"], errors="raise").dt.to_period("M")
    values = pd.to_numeric(data["cnr"], errors="raise")
    if periods.duplicated().any() or values.isna().any() or (values <= 0).any():
        raise ValueError("Les mois CNR doivent être uniques et leurs niveaux positifs.")
    return pd.Series(values.to_numpy(dtype=float), index=pd.PeriodIndex(periods), name="cnr").sort_index()


def load_ecb_eurusd(start: date, end: date) -> pd.Series:
    """Taux de référence BCE : dollars pour un euro, sans interpolation future."""
    url = f"{ECB_URL}?startPeriod={start.isoformat()}&endPeriod={end.isoformat()}&format=csvdata"
    request = Request(url, headers={"Accept": "text/csv", "User-Agent": "CNR-hedge-education/1.0"})
    with urlopen(request, timeout=15) as response:
        data = pd.read_csv(response)
    return parse_eurusd_csv(data)


def parse_eurusd_csv(data: pd.DataFrame) -> pd.Series:
    """Accepte un CSV simple ou l'export CSV officiel de la BCE."""
    date_column = "date" if "date" in data else "TIME_PERIOD"
    value_column = "eurusd" if "eurusd" in data else "OBS_VALUE"
    if date_column not in data or value_column not in data:
        raise ValueError("Le fichier EUR/USD doit contenir date,eurusd ou TIME_PERIOD,OBS_VALUE.")
    frame = pd.DataFrame(
        {
            "date": pd.to_datetime(data[date_column], errors="coerce").dt.normalize(),
            "eurusd": pd.to_numeric(data[value_column], errors="coerce"),
        }
    ).dropna()
    if frame.empty or frame["date"].duplicated().any() or (frame["eurusd"] <= 0).any():
        raise ValueError("Historique EUR/USD vide, non positif ou avec des dates répétées.")
    return frame.set_index("date")["eurusd"].sort_index()


def _contract_path(
    history: IceHistory, entry_date: pd.Timestamp, invoice_month: pd.Period
) -> tuple[pd.DataFrame, str, pd.Timestamp]:
    """Suit un seul ICE Brent de livraison M+2 malgré les changements de rang CO."""
    delivery = invoice_month + 2
    if delivery not in history.contracts.index:
        raise ValueError(f"Contrat ICE absent pour la facture {invoice_month}.")
    contract = history.contracts.loc[delivery]
    last_allowed = min(contract["expiry"], invoice_month.end_time.normalize())
    dates = history.prices.index[(history.prices.index >= entry_date) & (history.prices.index <= last_allowed)]
    if len(dates) < 2:
        raise ValueError(f"Pas assez de prix pour la facture {invoice_month}.")
    rows = []
    for quote_date in dates:
        active = history.contracts.index[history.contracts["expiry"] >= quote_date]
        if delivery not in active:
            raise ValueError(f"Échéance {delivery} inactive au {quote_date.date()}.")
        rank = int(active.get_loc(delivery)) + 1
        if rank > 12:
            raise ValueError(f"L'échéance {delivery} sort des douze rangs disponibles.")
        price = float(history.prices.at[quote_date, rank])
        rows.append((quote_date, rank, price))
    path = pd.DataFrame(rows, columns=["date", "co_rank", "price_usd_bbl"]).set_index("date")
    return path, str(contract["ticker"]), delivery.start_time


def _fx_on_dates(fx: float | pd.Series, dates: pd.DatetimeIndex) -> pd.Series:
    if isinstance(fx, (int, float)):
        rate = _number(fx, "EUR/USD")
        if rate <= 0:
            raise ValueError("Le taux EUR/USD doit être positif.")
        return pd.Series(rate, index=dates)
    reference = fx.copy()
    reference.index = pd.DatetimeIndex(reference.index).normalize()
    reference = reference.sort_index()
    if reference.index.has_duplicates or (reference <= 0).any():
        raise ValueError("Les dates EUR/USD doivent être uniques et les taux positifs.")
    aligned = reference.reindex(reference.index.union(dates)).sort_index().ffill().reindex(dates)
    if aligned.isna().any():
        raise ValueError("Il manque un taux EUR/USD connu au début de la couverture.")
    return aligned


def run_strip_backtest(
    history: IceHistory,
    cnr: pd.Series,
    *,
    entry_date: date,
    base_cnr_month: str | pd.Period,
    diesel_weight: float,
    positions: pd.DataFrame,
    eurusd: float | pd.Series,
    cnr_lag_months: int = 0,
) -> BacktestResult:
    """Calcule le P&L quotidien et le résidu mensuel pour chaque facture.

    positions : month (YYYY-MM), invoice_eur, barrels. La position reste sur le
    même future de livraison M+2 jusqu'à sa dernière cotation du mois M.
    """
    entry = pd.Timestamp(entry_date).normalize()
    if entry not in history.prices.index:
        raise ValueError("La date d'entrée doit être une séance présente dans le classeur.")
    if not 0 <= diesel_weight <= 1:
        raise ValueError("La pondération gazole doit être comprise entre 0 et 1.")
    if cnr_lag_months not in (0, 1, 2):
        raise ValueError("Le décalage CNR doit être de 0, 1 ou 2 mois.")
    base = pd.Period(base_cnr_month, freq="M")
    if base not in cnr.index:
        raise ValueError(f"Indice CNR de base absent : {base}.")
    required = {"month", "invoice_eur", "barrels"}
    if not required.issubset(positions.columns) or positions.empty:
        raise ValueError("Il faut des lignes month, invoice_eur et barrels.")
    frame = positions.copy()
    frame["month"] = pd.PeriodIndex(frame["month"], freq="M")
    if frame["month"].duplicated().any():
        raise ValueError("Chaque mois de facture doit apparaître une seule fois.")
    if frame["month"].min() < entry.to_period("M") or frame["month"].min() <= base:
        raise ValueError("Les factures doivent suivre la date d'entrée et le mois CNR de base.")
    frame["invoice_eur"] = pd.to_numeric(frame["invoice_eur"], errors="coerce")
    frame["barrels"] = pd.to_numeric(frame["barrels"], errors="coerce")
    if frame[["invoice_eur", "barrels"]].isna().any().any() or (frame[["invoice_eur", "barrels"]] < 0).any().any():
        raise ValueError("Les factures et le nombre de barils doivent être positifs ou nuls.")

    monthly_rows: list[dict[str, object]] = []
    daily_frames: list[pd.DataFrame] = []
    for position in frame.sort_values("month").itertuples(index=False):
        month = position.month
        cnr_month = month - cnr_lag_months
        if cnr_month not in cnr.index:
            raise ValueError(f"Indice CNR absent pour {cnr_month} (facture {month}).")
        path, ticker, delivery_date = _contract_path(history, entry, month)
        fx_rates = _fx_on_dates(eurusd, path.index)
        daily_usd = path["price_usd_bbl"].diff().fillna(0.0) * float(position.barrels)
        daily_eur = daily_usd / fx_rates
        path = path.assign(
            month=str(month),
            ticker=ticker,
            barrels=float(position.barrels),
            eurusd=fx_rates,
            pnl_usd_daily=daily_usd,
            pnl_eur_daily=daily_eur,
            pnl_eur_cum=daily_eur.cumsum(),
        )
        daily_frames.append(path.reset_index())
        surcharge = float(position.invoice_eur) * diesel_weight * (float(cnr.loc[cnr_month]) / float(cnr.loc[base]) - 1)
        pnl_eur = float(daily_eur.sum())
        front_entry = float(history.prices.at[path.index[0], 1])
        front_exit = float(history.prices.at[path.index[-1], 1])
        entry_spread = float(path["price_usd_bbl"].iloc[0]) - front_entry
        exit_spread = float(path["price_usd_bbl"].iloc[-1]) - front_exit
        barrels = float(position.barrels)
        monthly_rows.append(
            {
                "month": str(month),
                "cnr_month": str(cnr_month),
                "invoice_eur": float(position.invoice_eur),
                "cnr": float(cnr.loc[cnr_month]),
                "surcharge_eur": surcharge,
                "ticker": ticker,
                "delivery": delivery_date,
                "entry_date": path.index[0],
                "exit_date": path.index[-1],
                "rank_entry": int(path["co_rank"].iloc[0]),
                "rank_exit": int(path["co_rank"].iloc[-1]),
                "entry_price_usd_bbl": float(path["price_usd_bbl"].iloc[0]),
                "exit_price_usd_bbl": float(path["price_usd_bbl"].iloc[-1]),
                "term_spread_usd_bbl": entry_spread,
                "curve_shape": "contango" if entry_spread > 0 else "backwardation" if entry_spread < 0 else "plat",
                "front_move_usd": (front_exit - front_entry) * barrels,
                "curve_effect_usd": (exit_spread - entry_spread) * barrels,
                "barrels": barrels,
                "ice_contracts_fractional": barrels / CONTRACT_BARRELS_ICE,
                "mini_equivalents": barrels / CONTRACT_BARRELS_MINI,
                "pnl_usd": float(daily_usd.sum()),
                "pnl_eur": pnl_eur,
                "residual_eur": surcharge - pnl_eur,
            }
        )
    monthly = pd.DataFrame(monthly_rows)
    for name in ("surcharge_eur", "pnl_eur", "residual_eur"):
        monthly[f"{name}_cum"] = monthly[name].cumsum()
    daily = pd.concat(daily_frames, ignore_index=True)
    return BacktestResult(monthly, daily)
